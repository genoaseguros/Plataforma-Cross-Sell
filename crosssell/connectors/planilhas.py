"""Importação periódica de apólices de Saúde exportadas do Zeca.

Aceita CSV (; ou ,) e XLSX. Os cabeçalhos são mapeados via config/verticais.yaml,
então mudanças de layout da exportação não exigem mudança de código.

Cada importação é idempotente (chave: fonte + número da apólice). Apólices que
estavam ativas e não aparecem mais em uma importação *completa* do Zeca são
marcadas como canceladas — é assim que capturamos migrações entre operadoras.
"""

import csv
import io
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from crosssell.config import Settings
from crosssell.models import Negocio
from crosssell.normalize import normalizar_documento, parse_data, parse_numero, sem_acento
from crosssell.resolver import resolver_empresa

STATUS_INATIVO = ("cancel", "inativ", "suspens", "encerr", "rescind", "migr")


def _chave(texto) -> str:
    return " ".join(sem_acento(str(texto or "")).lower().replace("_", " ").split())


def ler_linhas(conteudo: bytes, nome_arquivo: str) -> list[dict]:
    if nome_arquivo.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        linhas = list(wb.active.iter_rows(values_only=True))
        if not linhas:
            return []
        cab = [_chave(c) for c in linhas[0]]
        return [dict(zip(cab, lin)) for lin in linhas[1:] if any(v not in (None, "") for v in lin)]
    texto = conteudo.decode("utf-8-sig", errors="replace")
    dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
    leitor = csv.DictReader(io.StringIO(texto), dialect=dialeto)
    return [{_chave(k): v for k, v in lin.items()} for lin in leitor]


def mapear(linha: dict, mapa: dict[str, list[str]]) -> dict:
    saida = {}
    for campo, aliases in mapa.items():
        for alias in aliases:
            v = linha.get(_chave(alias))
            if v not in (None, ""):
                saida[campo] = v
                break
    return saida


def _status(valor) -> str:
    t = _chave(valor)
    return "cancelado" if any(s in t for s in STATUS_INATIVO) else "ativo"


def importar_zeca(db: Session, settings: Settings, conteudo: bytes, nome_arquivo: str,
                  carga_completa: bool = True) -> dict:
    mapa = settings.verticais_config()["importacao"]["zeca"]
    vistos: set[str] = set()
    cont = {"linhas": 0, "importadas": 0, "sem_empresa": 0, "canceladas_por_ausencia": 0}

    for bruta in ler_linhas(conteudo, nome_arquivo):
        cont["linhas"] += 1
        r = mapear(bruta, mapa)
        cnpj, _ = normalizar_documento(r.get("cnpj"))
        empresa = resolver_empresa(db, razao_social=r.get("razao_social"), cnpj=cnpj)
        if empresa is None:
            cont["sem_empresa"] += 1
            continue
        numero = str(r.get("numero_apolice") or f"{empresa.cnpj or empresa.id}-{r.get('operadora', '')}")
        vistos.add(numero)
        neg = db.scalar(select(Negocio).where(Negocio.fonte == "zeca", Negocio.id_externo == numero))
        if neg is None:
            neg = Negocio(fonte="zeca", id_externo=numero, vertical="saude", status="ativo")
            db.add(neg)
        neg.empresa = empresa
        neg.titulo = f"Saúde {r.get('operadora', '')}".strip()
        neg.produto = r.get("produto")
        neg.seguradora = r.get("operadora")
        neg.inicio_vigencia = parse_data(r.get("inicio_vigencia"))
        neg.fim_vigencia = parse_data(r.get("fim_vigencia"))
        neg.valor = parse_numero(r.get("premio"))
        vidas = parse_numero(r.get("vidas"))
        neg.vidas = int(vidas) if vidas is not None else None
        neg.status = _status(r.get("status"))
        cont["importadas"] += 1

    if carga_completa and vistos:
        for neg in db.scalars(select(Negocio).where(Negocio.fonte == "zeca", Negocio.status == "ativo")):
            if neg.id_externo not in vistos:
                neg.status = "cancelado"
                cont["canceladas_por_ausencia"] += 1

    db.commit()
    return cont


def importar_arquivo(db: Session, settings: Settings, fonte: str, caminho: Path, **kw) -> dict:
    fn = {"zeca": importar_zeca}[fonte]
    return fn(db, settings, caminho.read_bytes(), caminho.name, **kw)
